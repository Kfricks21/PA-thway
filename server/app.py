from pathlib import Path
import json
import os
import re
import sqlite3
import tempfile
from urllib import error as urllib_error
from urllib import request as urllib_request

from flask import Flask, jsonify, request, send_from_directory
from pptx import Presentation
from pypdf import PdfReader

from hub import register_hub
from pathway_app import register_pathway

BASE_DIR = Path(__file__).resolve().parent.parent
CLIENT_DIR = BASE_DIR / "client" / "dist"
DB_PATH = BASE_DIR / "server" / "storage" / "student_hub.db"

app = Flask(__name__)

# The PA-thway student app at /app, and the shared Student Hub at /api/hub/*.
register_pathway(app)
register_hub(app)


def get_db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_student_hub_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = get_db_connection()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            cohort TEXT NOT NULL,
            text TEXT NOT NULL,
            likes INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS note_replies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            note_id INTEGER NOT NULL,
            author TEXT NOT NULL,
            message TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(note_id) REFERENCES notes(id)
        );

        CREATE TABLE IF NOT EXISTS comments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            author TEXT NOT NULL,
            cohort TEXT NOT NULL,
            message TEXT NOT NULL,
            likes INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS comment_replies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            comment_id INTEGER NOT NULL,
            author TEXT NOT NULL,
            message TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(comment_id) REFERENCES comments(id)
        );

        CREATE TABLE IF NOT EXISTS challenges (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cohort TEXT NOT NULL,
            target_cohort TEXT NOT NULL,
            prompt TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        """
    )

    conn.execute(
        "DELETE FROM notes WHERE title IN ('Course recap', 'Student tips')"
    )
    conn.execute(
        "DELETE FROM comments WHERE author IN ('Jordan', 'Alicia')"
    )
    conn.execute(
        "DELETE FROM challenges WHERE prompt IN ("
        "'Create a 3-question follow-up quiz from this week’s slides.', "
        "'Identify the strongest question style for student review.', "
        "'Turn one topic into a challenge activity for peers.')"
    )
    conn.commit()
    conn.close()


def serialize_note(row, replies):
    return {
        "id": row["id"],
        "title": row["title"],
        "cohort": row["cohort"],
        "text": row["text"],
        "likes": row["likes"],
        "replies": replies,
    }


def serialize_comment(row, replies):
    return {
        "id": row["id"],
        "author": row["author"],
        "cohort": row["cohort"],
        "message": row["message"],
        "likes": row["likes"],
        "replies": replies,
    }


def serialize_challenge(row):
    return {
        "id": row["id"],
        "cohort": row["cohort"],
        "targetCohort": row["target_cohort"],
        "prompt": row["prompt"],
    }


def fetch_hub_state():
    conn = get_db_connection()
    notes = []
    for note_row in conn.execute("SELECT * FROM notes ORDER BY id DESC"):
        note_replies = [
            {
                "id": reply_row["id"],
                "author": reply_row["author"],
                "message": reply_row["message"],
            }
            for reply_row in conn.execute(
                "SELECT * FROM note_replies WHERE note_id = ? ORDER BY id ASC",
                (note_row["id"],),
            )
        ]
        notes.append(serialize_note(note_row, note_replies))

    comments = []
    for comment_row in conn.execute("SELECT * FROM comments ORDER BY id DESC"):
        comment_replies = [
            {
                "id": reply_row["id"],
                "author": reply_row["author"],
                "message": reply_row["message"],
            }
            for reply_row in conn.execute(
                "SELECT * FROM comment_replies WHERE comment_id = ? ORDER BY id ASC",
                (comment_row["id"],),
            )
        ]
        comments.append(serialize_comment(comment_row, comment_replies))

    challenges = [serialize_challenge(row) for row in conn.execute("SELECT * FROM challenges ORDER BY id DESC")]
    conn.close()

    return {
        "comments": comments,
        "notes": notes,
        "challenges": challenges,
        "cohortCode": "PA-101",
    }


init_student_hub_db()


def extract_content_details(uploaded_file):
    if uploaded_file is None or not uploaded_file.filename:
        return {"text": "", "segments": []}

    suffix = Path(uploaded_file.filename).suffix.lower()
    fd, temp_path = tempfile.mkstemp(suffix=suffix)
    os.close(fd)

    try:
        uploaded_file.save(temp_path)

        if suffix == ".pptx":
            presentation = Presentation(temp_path)
            segments = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                blocks = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        cleaned = shape.text.strip()
                        if cleaned:
                            blocks.append(cleaned)

                slide_text = "\n".join(block for block in blocks if block)
                if slide_text:
                    segments.append(
                        {
                            "type": "slide",
                            "number": slide_number,
                            "text": slide_text,
                            "blocks": blocks,
                        }
                    )

            text = "\n".join(segment["text"] for segment in segments)
            return {"text": text, "segments": segments}

        if suffix == ".pdf":
            reader = PdfReader(temp_path)
            segments = []
            for page_number, page in enumerate(reader.pages, start=1):
                content = (page.extract_text() or "").strip()
                if content:
                    segments.append(
                        {
                            "type": "page",
                            "number": page_number,
                            "text": content,
                        }
                    )

            text = "\n".join(segment["text"] for segment in segments)
            return {"text": text, "segments": segments}

        if suffix in {".txt", ".md"}:
            with open(temp_path, "r", encoding="utf-8", errors="ignore") as handle:
                content = handle.read()

            return {
                "text": content,
                "segments": [{"type": "document", "number": 1, "text": content}],
            }

        return {"text": "", "segments": []}
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def extract_text_from_upload(uploaded_file):
    return extract_content_details(uploaded_file).get("text", "")


def infer_domain(content):
    content_lower = content.lower()
    medical_terms = ["patient", "diagnosis", "symptom", "treatment", "disease", "therapy", "clinical", "medication", "outcome"]
    science_terms = ["cell", "molecule", "hypothesis", "experiment", "gene", "lab", "process", "mechanism", "signal", "species"]

    medical_hits = sum(1 for term in medical_terms if term in content_lower)
    science_hits = sum(1 for term in science_terms if term in content_lower)

    if medical_hits >= science_hits:
        return "medical"

    return "scientific" if science_hits else "general"


def normalize_question_item(question, idx):
    prompt = question.get("prompt") or f"Question {idx + 1}: Based on the uploaded slide content, what is the best answer?"
    options = question.get("options") or []

    if len(options) < 4:
        options = [
            question.get("answer") or "The slide supports this statement.",
            "A related but incomplete interpretation.",
            "A statement that is not supported by the uploaded material.",
            "A generic answer that does not match the slide content.",
        ]

    answer = question.get("answer") or options[0]
    level = question.get("level") or ("second_order" if idx % 2 else "first_order")

    return {
        "id": idx + 1,
        "type": "multiple-choice",
        "level": level,
        "prompt": prompt,
        "options": options[:4],
        "answer": answer,
    }


def parse_openai_response(response_text):
    cleaned = response_text.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        if cleaned.startswith("json"):
            cleaned = cleaned[4:].strip()

    return json.loads(cleaned)


def generate_questions_with_openai(content, filename, question_count):
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        return None

    api_url = os.getenv("OPENAI_API_URL", "https://api.openai.com/v1/chat/completions")
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")

    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are an educational content expert. Generate a quiz from the uploaded slide text. "
                    "Create first-order (recall/definition) and second-order (application/interpretation) questions. "
                    "Make the questions scientific or medical in tone and based only on the slide content. "
                    "Return valid JSON only with an array of question objects containing id, level, prompt, options, and answer."
                ),
            },
            {
                "role": "user",
                "content": (
                    f"Generate exactly {question_count} questions based on this slide content from {filename}. "
                    f"The questions should be scientific or medical and should mix first-order and second-order questions.\n\n"
                    f"Slide content:\n{content[:12000]}"
                ),
            },
        ],
        "temperature": 0.7,
    }

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    try:
        req = urllib_request.Request(api_url, data=json.dumps(payload).encode("utf-8"), headers=headers, method="POST")
        with urllib_request.urlopen(req, timeout=30) as response:
            response_payload = json.loads(response.read().decode("utf-8"))

        raw_content = response_payload["choices"][0]["message"]["content"]
        parsed = parse_openai_response(raw_content)

        if isinstance(parsed, list) and parsed:
            normalized = [normalize_question_item(item, idx + 1) for idx, item in enumerate(parsed[:question_count])]
            if normalized:
                return normalized
    except (urllib_error.HTTPError, urllib_error.URLError, ValueError, KeyError, json.JSONDecodeError):
        return None

    return None


def build_interactive_modules(content, filename):
    cleaned_content = re.sub(r"\s+", " ", content or "").strip()
    if not cleaned_content:
        cleaned_content = f"Uploaded presentation: {filename}"

    content_lower = cleaned_content.lower()
    has_action_potential = any(
        keyword in content_lower
        for keyword in [
            "action potential",
            "membrane potential",
            "depolarization",
            "repolarization",
            "threshold potential",
            "na+",
            "sodium",
            "potassium",
            "calcium",
            "chloride",
        ]
    )

    if not has_action_potential:
        return []

    return [
        {
            "id": "action-potential-explorer",
            "title": "Action Potential Explorer",
            "type": "action-potential",
            "description": "Adjust ion inputs to move the membrane potential through a slide-aligned action potential cycle.",
            "context": f"Derived from {filename}: uploaded slide content points toward ion movement and membrane potential changes.",
            "starterPotential": -70,
            "ions": [
                {"label": "Na+", "effect": 24, "description": "Fast sodium influx drives depolarization"},
                {"label": "K+", "effect": -18, "description": "Potassium efflux supports repolarization"},
                {"label": "Ca2+", "effect": 12, "description": "Calcium supports signal strength"},
                {"label": "Cl-", "effect": -10, "description": "Chloride can dampen or stabilize the signal"},
            ],
        }
    ]


def build_questions_from_content(content_details, filename, question_count):
    if isinstance(content_details, dict):
        cleaned_content = re.sub(r"\s+", " ", content_details.get("text") or "").strip()
        segments = content_details.get("segments") or []
    else:
        cleaned_content = re.sub(r"\s+", " ", content_details or "").strip()
        segments = []

    if not cleaned_content:
        cleaned_content = f"Uploaded presentation: {filename}"

    if not segments:
        segments = [{"type": "document", "number": 1, "text": cleaned_content}]

    ai_questions = generate_questions_with_openai(cleaned_content, filename, question_count)
    if ai_questions:
        for idx, question in enumerate(ai_questions):
            source = segments[idx % len(segments)]
            question["source"] = {
                "type": source.get("type", "document"),
                "number": source.get("number", 1),
                "text": source.get("text", cleaned_content)[:200].strip(),
            }
        return ai_questions

    sentences = [part.strip() for part in re.split(r"(?<=[.!?])\s+", cleaned_content) if part.strip()]
    if not sentences:
        sentences = [cleaned_content]

    domain = infer_domain(cleaned_content)
    questions = []

    for idx in range(question_count):
        source = segments[idx % len(segments)]
        source_text = source.get("text") or cleaned_content
        snippet = source_text[:180].rstrip(". ") or source_text[:180]

        if idx % 2 == 0:
            answer = snippet
            prompt = (
                f"Question {idx + 1}: Based on the uploaded {domain} slide content, which statement most accurately reflects the key point described?"
            )
            options = [
                answer,
                f"A related but inaccurate interpretation of the {domain} content.",
                f"A statement that is not supported by the uploaded slide material.",
                f"A generic or unrelated response that does not match the presentation."
            ]
            level = "first_order"
        else:
            answer = (
                f"The slide content supports the interpretation that {snippet.lower()} has important {domain} relevance and should be considered in context."
            )
            prompt = (
                f"Question {idx + 1}: Based on the uploaded {domain} slide content, which interpretation is best supported by the information presented?"
            )
            options = [
                answer,
                f"The topic is unrelated to the evidence shown on the slide and should be treated as separate from the {domain} discussion.",
                f"The conclusion is based on assumptions that are not stated or supported by the uploaded material.",
                f"The content suggests a conclusion that contradicts the evidence from the slides."
            ]
            level = "second_order"

        questions.append(
            {
                "id": idx + 1,
                "type": "multiple-choice",
                "level": level,
                "prompt": prompt,
                "options": options[:4],
                "answer": answer,
                "source": {
                    "type": source.get("type", "document"),
                    "number": source.get("number", 1),
                    "text": snippet,
                },
            }
        )

    return questions


@app.get("/api/health")
def health():
    return jsonify(
        {
            "status": "ok",
            "service": "ppt-quiz-generator",
            "client_build_exists": CLIENT_DIR.exists(),
        }
    )


@app.get("/api/hub/state")
def get_hub_state_api():
    return jsonify(fetch_hub_state())


@app.post("/api/hub/notes")
def create_note_api():
    payload = request.get_json(silent=True) or {}
    title = (payload.get("title") or "Shared note").strip() or "Shared note"
    cohort = (payload.get("cohort") or "General").strip() or "General"
    text = (payload.get("text") or "").strip()

    if not text:
        return jsonify({"error": "Note text is required."}), 400

    conn = get_db_connection()
    cursor = conn.execute(
        "INSERT INTO notes (title, cohort, text, likes) VALUES (?, ?, ?, 0)",
        (title, cohort, text),
    )
    conn.commit()

    note = conn.execute("SELECT * FROM notes WHERE id = ?", (cursor.lastrowid,)).fetchone()
    conn.close()

    return jsonify({"note": serialize_note(note, [])}), 201


@app.post("/api/hub/comments")
def create_comment_api():
    payload = request.get_json(silent=True) or {}
    author = (payload.get("author") or "You").strip() or "You"
    cohort = (payload.get("cohort") or "General").strip() or "General"
    message = (payload.get("message") or "").strip()

    if not message:
        return jsonify({"error": "Comment text is required."}), 400

    conn = get_db_connection()
    cursor = conn.execute(
        "INSERT INTO comments (author, cohort, message, likes) VALUES (?, ?, ?, 0)",
        (author, cohort, message),
    )
    conn.commit()

    comment = conn.execute("SELECT * FROM comments WHERE id = ?", (cursor.lastrowid,)).fetchone()
    conn.close()

    return jsonify({"comment": serialize_comment(comment, [])}), 201


@app.post("/api/hub/challenges")
def create_challenge_api():
    payload = request.get_json(silent=True) or {}
    cohort = (payload.get("cohort") or "General").strip() or "General"
    target_cohort = (payload.get("targetCohort") or cohort).strip() or cohort
    prompt = (payload.get("prompt") or "").strip()

    if not prompt:
        return jsonify({"error": "Challenge prompt is required."}), 400

    conn = get_db_connection()
    cursor = conn.execute(
        "INSERT INTO challenges (cohort, target_cohort, prompt) VALUES (?, ?, ?)",
        (cohort, target_cohort, prompt),
    )
    conn.commit()

    challenge = conn.execute("SELECT * FROM challenges WHERE id = ?", (cursor.lastrowid,)).fetchone()
    conn.close()

    return jsonify({"challenge": serialize_challenge(challenge)}), 201


@app.post("/api/hub/replies")
def create_reply_api():
    payload = request.get_json(silent=True) or {}
    target_type = payload.get("type")
    target_id = payload.get("targetId")
    author = (payload.get("author") or "You").strip() or "You"
    message = (payload.get("message") or "").strip()

    if target_type not in {"note", "comment"} or not target_id:
        return jsonify({"error": "Invalid reply target."}), 400

    if not message:
        return jsonify({"error": "Reply text is required."}), 400

    conn = get_db_connection()

    if target_type == "note":
        conn.execute(
            "INSERT INTO note_replies (note_id, author, message) VALUES (?, ?, ?)",
            (target_id, author, message),
        )
        reply = conn.execute(
            "SELECT id, author, message FROM note_replies WHERE note_id = ? ORDER BY id DESC LIMIT 1",
            (target_id,),
        ).fetchone()
    else:
        conn.execute(
            "INSERT INTO comment_replies (comment_id, author, message) VALUES (?, ?, ?)",
            (target_id, author, message),
        )
        reply = conn.execute(
            "SELECT id, author, message FROM comment_replies WHERE comment_id = ? ORDER BY id DESC LIMIT 1",
            (target_id,),
        ).fetchone()

    conn.commit()
    conn.close()

    return jsonify({"reply": {"id": reply["id"], "author": reply["author"], "message": reply["message"]}}), 201


@app.post("/api/hub/likes")
def toggle_like_api():
    payload = request.get_json(silent=True) or {}
    target_type = payload.get("type")
    target_id = payload.get("targetId")

    if target_type not in {"note", "comment"} or not target_id:
        return jsonify({"error": "Invalid like target."}), 400

    conn = get_db_connection()

    if target_type == "note":
        conn.execute("UPDATE notes SET likes = likes + 1 WHERE id = ?", (target_id,))
        row = conn.execute("SELECT likes FROM notes WHERE id = ?", (target_id,)).fetchone()
    else:
        conn.execute("UPDATE comments SET likes = likes + 1 WHERE id = ?", (target_id,))
        row = conn.execute("SELECT likes FROM comments WHERE id = ?", (target_id,)).fetchone()

    conn.commit()
    conn.close()

    return jsonify({"likes": row["likes"]})


@app.post("/api/generate")
def generate_quiz():
    uploaded_file = request.files.get("file")
    filename = uploaded_file.filename if uploaded_file and uploaded_file.filename else request.form.get("filename") or "presentation.pptx"
    question_count = request.form.get("questionCount", default=5, type=int)

    if question_count <= 0:
        question_count = 5

    extracted_content = extract_content_details(uploaded_file)
    questions = build_questions_from_content(extracted_content, filename, question_count)
    interactive_modules = build_interactive_modules(extracted_content, filename)

    return jsonify(
        {
            "status": "success",
            "message": "Quiz generation request received.",
            "filename": filename,
            "questions": questions,
            "interactiveModules": interactive_modules,
        }
    )


@app.route("/", defaults={"path": ""})
@app.route("/<path:path>")
def serve_frontend(path):
    if not CLIENT_DIR.exists():
        return jsonify(
            {
                "error": "Frontend build not found. Run npm --prefix client run build before starting the server."
            }
        ), 404

    if path and (CLIENT_DIR / path).exists() and (CLIENT_DIR / path).is_file():
        return send_from_directory(str(CLIENT_DIR), path)

    return send_from_directory(str(CLIENT_DIR), "index.html")


if __name__ == "__main__":
    host = os.getenv("HOST", "0.0.0.0")
    port = int(os.getenv("PORT", "5000"))
    app.run(host=host, port=port, debug=False)
